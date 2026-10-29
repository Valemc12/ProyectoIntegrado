"""
Base genérica de los mantenedores (CRUD) del SGR.

Cada mantenedor define su modelo, formulario, columnas, campos de búsqueda,
filtros y su alcance por delegación. Esta clase entrega las vistas comunes:
listar (búsqueda + filtros + paginador), consultar, crear, editar, eliminar,
activar/desactivar y exportar a Excel. Todas validan permisos en el servidor.
"""
import logging
from dataclasses import dataclass
from datetime import datetime
from functools import reduce
from io import BytesIO
from operator import or_
from typing import Callable

from django import forms
from django.contrib import messages
from django.core.paginator import Paginator
from django.db import DatabaseError, transaction
from django.db.models import ProtectedError, Q
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from django.urls import path, reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .permisos import NOMBRE_MODULO, requiere_permiso

log = logging.getLogger("sgr")

ROJO_HERALDICO = "C41230"


def obtener_valor(obj, accesor):
    if callable(accesor):
        return accesor(obj)
    valor = obj
    for parte in accesor.split("."):
        valor = getattr(valor, parte, None)
        if valor is None:
            return ""
    return valor() if callable(valor) else valor


@dataclass
class Columna:
    titulo: str
    accesor: str | Callable
    tipo: str = "texto"  # texto | estado | enlace


@dataclass
class Filtro:
    nombre: str
    etiqueta: str
    opciones: Callable  # (request) -> [(valor, texto)]
    aplicar: Callable  # (queryset, valor) -> queryset


def filtro_estado(campo="activo", femenino=False):
    act, inact = ("Activa", "Inactiva") if femenino else ("Activo", "Inactivo")
    return Filtro(
        "estado", "Estado",
        lambda request: [("activo", act), ("inactivo", inact)],
        lambda qs, v: qs.filter(**{campo: v == "activo"}) if v in ("activo", "inactivo") else qs,
    )


class FormularioBase(forms.ModelForm):
    """Normaliza textos, valida duplicados sin distinguir mayúsculas y aplica estilos."""

    unicos = {"codigo": "Ya existe un registro con este código."}
    minimos = {"nombre": 3}

    def __init__(self, *args, request=None, **kwargs):
        self.request = request
        super().__init__(*args, **kwargs)
        for nombre, campo in self.fields.items():
            widget = campo.widget
            if isinstance(widget, (forms.CheckboxSelectMultiple, forms.RadioSelect)):
                continue
            if isinstance(widget, forms.CheckboxInput):
                widget.attrs.setdefault("class", "check")
                continue
            clases = "campo-select" if isinstance(widget, forms.Select) else "campo-input"
            widget.attrs.setdefault("class", clases)
            if campo.required:
                widget.attrs["required"] = True
            if getattr(campo, "max_length", None):
                widget.attrs.setdefault("maxlength", campo.max_length)
            if nombre in self.minimos:
                widget.attrs.setdefault("minlength", self.minimos[nombre])
            if isinstance(widget, forms.Textarea):
                widget.attrs.setdefault("rows", 3)
        if "codigo" in self.fields:
            self.fields["codigo"].widget.attrs.update(
                {"pattern": "[A-Za-z0-9][A-Za-z0-9-]{1,14}", "style": "text-transform:uppercase",
                 "placeholder": "Ej.: COD01"}
            )

    def clean(self):
        datos = super().clean()
        for nombre, valor in list(datos.items()):
            if isinstance(valor, str) and not nombre.startswith("clave"):
                datos[nombre] = " ".join(valor.split()) if nombre != "descripcion" else valor.strip()
        if datos.get("codigo"):
            datos["codigo"] = datos["codigo"].upper()
        for nombre, minimo in self.minimos.items():
            valor = datos.get(nombre)
            if isinstance(valor, str) and valor and len(valor) < minimo:
                self.add_error(nombre, f"Debe tener al menos {minimo} caracteres.")
        modelo = self._meta.model
        for nombre, mensaje in self.unicos.items():
            valor = datos.get(nombre)
            if not valor or nombre in self.errors:
                continue
            qs = modelo.objects.filter(**{f"{nombre}__iexact": valor})
            if self.instance.pk:
                qs = qs.exclude(pk=self.instance.pk)
            if qs.exists():
                self.add_error(nombre, mensaje)
        return datos

    def validate_unique(self):
        pass


class Mantenedor:
    modulo = ""
    model = None
    form_class = None
    titulo = ""
    singular = ""
    femenino = False
    campos_busqueda: list = []
    columnas: list = []
    columnas_excel: list = []
    detalle: list = []
    filtros: list = []
    campo_estado = "activo"
    por_pagina = 10
    plantilla_lista = "crud/lista.html"
    plantilla_form = "crud/form.html"
    plantilla_detalle = "crud/detalle.html"
    placeholder_busqueda = "Buscar por código, nombre u otro campo…"
    ayuda_lista = ""

    def queryset(self, request):
        return self.model.objects.all()

    def motivo_no_eliminar(self, request, obj):
        return None

    def motivo_no_cambiar_estado(self, request, obj):
        return None

    def form_kwargs(self, request, obj=None):
        return {"request": request}

    def guardar(self, request, form, creando):
        return form.save()

    def contexto_extra(self, request, obj=None):
        return {}

    def acciones_extra(self, request, obj):
        return []

    @property
    def nombre_modulo(self):
        return NOMBRE_MODULO.get(self.modulo, self.titulo)

    def url(self, nombre, *args):
        return reverse(f"{self.modulo}:{nombre}", args=args)

    def obtener(self, request, pk):
        try:
            return self.queryset(request).get(pk=pk)
        except (self.model.DoesNotExist, ValueError, TypeError):
            raise Http404("Registro no encontrado o fuera de su alcance autorizado.")

    def filtrar(self, request):
        qs = self.queryset(request)
        q = " ".join(request.GET.get("q", "").split())[:100]
        if q:
            qs = qs.filter(reduce(or_, (Q(**{f"{c}__icontains": q}) for c in self.campos_busqueda)))
        valores = {}
        for filtro in self.filtros:
            valor = request.GET.get(filtro.nombre, "")
            validos = {str(v) for v, _ in filtro.opciones(request)}
            if valor and valor in validos:
                qs = filtro.aplicar(qs, valor)
                valores[filtro.nombre] = valor
        return qs.distinct(), q, valores

    def es_estado(self, obj):
        return getattr(obj, self.campo_estado)

    def filas(self, request, objetos):
        filas = []
        for obj in objetos:
            celdas = []
            for col in self.columnas:
                valor = obtener_valor(obj, col.accesor)
                celdas.append({"tipo": col.tipo, "valor": valor})
            filas.append({"obj": obj, "celdas": celdas, "activo": self.es_estado(obj) if self.campo_estado else None,
                          "extras": self.acciones_extra(request, obj)})
        return filas

    def contexto_base(self, request, **extra):
        ctx = {
            "m": self,
            "modulo": self.modulo,
            "titulo": self.titulo,
            "singular": self.singular,
            "nuevo_texto": ("Nueva " if self.femenino else "Nuevo ") + self.singular,
            "puede_crear": request.user.tiene_permiso(f"{self.modulo}.crear"),
            "puede_editar": request.user.tiene_permiso(f"{self.modulo}.editar"),
            "puede_eliminar": request.user.tiene_permiso(f"{self.modulo}.eliminar"),
            "puede_exportar": request.user.tiene_permiso(f"{self.modulo}.exportar"),
        }
        ctx.update(extra)
        return ctx

    # --- Vistas -----------------------------------------------------------
    def vista_lista(self, request):
        qs, q, valores = self.filtrar(request)
        
        # Paginación configurable [10 | 20 | 50 | 100]
        limite_raw = request.GET.get("limite")
        por_pagina = self.por_pagina
        if limite_raw and limite_raw.isdigit() and int(limite_raw) in (10, 20, 50, 100):
            por_pagina = int(limite_raw)

        paginador = Paginator(qs, por_pagina)
        pagina = paginador.get_page(request.GET.get("pagina"))
        parametros = request.GET.copy()
        parametros.pop("pagina", None)
        filtros = [
            {"nombre": f.nombre, "etiqueta": f.etiqueta, "opciones": f.opciones(request),
             "valor": valores.get(f.nombre, "")}
            for f in self.filtros
        ]
        ctx = self.contexto_base(
            request,
            pagina=pagina,
            filas=self.filas(request, pagina.object_list),
            total=paginador.count,
            hay_registros=self.queryset(request).exists(),
            q=q,
            filtros=filtros,
            filtrado=bool(q or valores),
            parametros=parametros.urlencode(),
            columnas=self.columnas,
            limite_actual=por_pagina,
        )
        ctx.update(self.contexto_extra(request))
        return render(request, self.plantilla_lista, ctx)

    def vista_detalle(self, request, pk):
        obj = self.obtener(request, pk)
        datos = []
        for etiqueta, accesor in self.detalle:
            datos.append((etiqueta, obtener_valor(obj, accesor)))
        ctx = self.contexto_base(request, obj=obj, datos=datos,
                                 motivo_no_eliminar=self.motivo_no_eliminar(request, obj),
                                 activo=self.es_estado(obj) if self.campo_estado else None)
        ctx.update(self.contexto_extra(request, obj))
        return render(request, self.plantilla_detalle, ctx)

    def _formulario(self, request, obj=None):
        creando = obj is None
        form = self.form_class(request.POST if request.method == "POST" else None, instance=obj, **self.form_kwargs(request, obj))
        if request.method == "POST":
            if form.is_valid():
                try:
                    with transaction.atomic():
                        guardado = self.guardar(request, form, creando)
                except DatabaseError:
                    log.exception("Error al guardar %s", self.modulo)
                    messages.error(request, "No fue posible guardar por un error del servidor. "
                                            "Intente nuevamente; no se registraron cambios.")
                else:
                    accion = "creado" if creando else "actualizado"
                    if self.femenino:
                        accion = accion[:-1] + "a"
                    messages.success(request, f"{self.singular.capitalize()} «{guardado}» {accion} correctamente.")
                    return redirect(self.url("detalle", guardado.pk))
        ctx = self.contexto_base(
            request, form=form, obj=obj, creando=creando,
            errores_formulario=request.method == "POST" and bool(form.errors),
        )
        ctx.update(self.contexto_extra(request, obj))
        return render(request, self.plantilla_form, ctx)

    def vista_crear(self, request):
        return self._formulario(request)

    def vista_editar(self, request, pk):
        return self._formulario(request, self.obtener(request, pk))

    def vista_eliminar(self, request, pk):
        obj = self.obtener(request, pk)
        motivo = self.motivo_no_eliminar(request, obj)
        if motivo:
            messages.error(request, motivo)
            return redirect(self.url("detalle", obj.pk))
        texto = str(obj)
        try:
            with transaction.atomic():
                obj.delete()
        except ProtectedError:
            messages.error(request, f"No se puede eliminar «{texto}» porque tiene registros relacionados. "
                                    "Puede desactivarlo en su lugar.")
            return redirect(self.url("detalle", obj.pk))
        except DatabaseError:
            log.exception("Error al eliminar %s", self.modulo)
            messages.error(request, "No fue posible eliminar por un error del servidor. No se realizaron cambios.")
            return redirect(self.url("detalle", obj.pk))
        accion = "eliminada" if self.femenino else "eliminado"
        messages.success(request, f"{self.singular.capitalize()} «{texto}» {accion} correctamente.")
        return redirect(self.url("lista"))

    def vista_cambiar_estado(self, request, pk):
        obj = self.obtener(request, pk)
        motivo = self.motivo_no_cambiar_estado(request, obj)
        if motivo:
            messages.error(request, motivo)
        else:
            nuevo = not getattr(obj, self.campo_estado)
            setattr(obj, self.campo_estado, nuevo)
            obj.save(update_fields=[self.campo_estado, "actualizado"])
            estado = obj.estado_texto.lower()
            messages.success(request, f"{self.singular.capitalize()} «{obj}» ahora está {estado}.")
        destino = request.POST.get("volver")
        return redirect(self.url("lista") + (f"?{destino}" if destino and "//" not in destino else ""))

    def vista_exportar(self, request):
        qs, q, valores = self.filtrar(request)
        libro = Workbook()
        hoja = libro.active
        hoja.title = self.nombre_modulo[:31]
        columnas = self.columnas_excel or [(c.titulo, c.accesor) for c in self.columnas]
        hoja.append([t for t, _ in columnas])
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=ROJO_HERALDICO)
            celda.alignment = Alignment(vertical="center")
        for obj in qs.iterator(chunk_size=500) if not qs._prefetch_related_lookups else qs:
            fila = []
            for _, accesor in columnas:
                valor = obtener_valor(obj, accesor)
                if isinstance(valor, bool):
                    valor = "Activo" if valor else "Inactivo"
                elif isinstance(valor, datetime):
                    valor = timezone.localtime(valor).replace(tzinfo=None)
                elif not isinstance(valor, (int, float)):
                    valor = str(valor)
                fila.append(valor)
            hoja.append(fila)
        for i, (titulo, _) in enumerate(columnas, start=1):
            ancho = max([len(str(titulo))] + [len(str(c.value or "")) for c in hoja[get_column_letter(i)][1:200]])
            hoja.column_dimensions[get_column_letter(i)].width = min(60, ancho + 3)
        hoja.auto_filter.ref = hoja.dimensions
        hoja.freeze_panes = "A2"
        info = libro.create_sheet("Filtro aplicado")
        info.append(["Módulo", self.nombre_modulo])
        info.append(["Delegación activa", str(request.delegacion_activa)])
        info.append(["Búsqueda", q or "(sin búsqueda)"])
        for nombre, valor in valores.items():
            info.append([f"Filtro {nombre}", valor])
        info.append(["Total de registros", qs.count()])
        info.append(["Exportado por", f"{request.user.nombre} ({request.user.rol})"])
        info.append(["Fecha", timezone.localtime().strftime("%d-%m-%Y %H:%M")])
        salida = BytesIO()
        libro.save(salida)
        nombre = f"{self.modulo}_{timezone.localtime().strftime('%Y%m%d_%H%M')}.xlsx"
        respuesta = HttpResponse(
            salida.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        respuesta["Content-Disposition"] = f'attachment; filename="{nombre}"'
        return respuesta

    def urls(self):
        p = self.modulo

        def con(permiso, vista, solo_post=False):
            v = requiere_permiso(f"{p}.ver", f"{p}.{permiso}")(vista) if permiso != "ver" else \
                requiere_permiso(f"{p}.ver")(vista)
            return require_POST(v) if solo_post else v

        patrones = [
            path("", con("ver", self.vista_lista), name="lista"),
            path("nuevo/", con("crear", self.vista_crear), name="crear"),
            path("exportar/", con("exportar", self.vista_exportar), name="exportar"),
            path("<int:pk>/", con("ver", self.vista_detalle), name="detalle"),
            path("<int:pk>/editar/", con("editar", self.vista_editar), name="editar"),
            path("<int:pk>/eliminar/", con("eliminar", self.vista_eliminar, True), name="eliminar"),
            path("<int:pk>/estado/", con("eliminar", self.vista_cambiar_estado, True), name="estado"),
        ]
        return patrones, p